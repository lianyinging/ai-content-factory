#!/usr/bin/env python3
"""Import 展厅门店场景产品清单.xlsx into resources/products/."""
from __future__ import annotations

import argparse
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "resources" / "products"


def clean_text(v) -> str:
    if v is None:
        return ""
    return str(v).strip().replace("\n", "").replace("\r", "")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--max-side", type=int, default=1200)
    ap.add_argument("--quality", type=int, default=85)
    args = ap.parse_args()

    out: Path = args.out
    img_dir = out / "images"
    img_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(args.xlsx, data_only=True)
    ws = wb.active

    row_images: dict[int, tuple[str | None, bytes]] = {}
    for img in ws._images:
        row = img.anchor._from.row + 1
        row_images[row] = (getattr(img, "format", None), img._data())

    products = []
    for r in range(2, ws.max_row + 1):
        sku = clean_text(ws.cell(r, 1).value)
        name = clean_text(ws.cell(r, 2).value)
        if not sku or not name:
            continue
        if r not in row_images:
            print(f"skip {sku}: no image")
            continue
        _fmt, raw = row_images[r]
        im = Image.open(io.BytesIO(raw))
        if im.mode in ("RGBA", "P", "LA"):
            bg = Image.new("RGB", im.size, (255, 255, 255))
            if im.mode == "P":
                im = im.convert("RGBA")
            bg.paste(im, mask=im.split()[-1] if im.mode in ("RGBA", "LA") else None)
            im = bg
        elif im.mode != "RGB":
            im = im.convert("RGB")
        w, h = im.size
        if max(w, h) > args.max_side:
            scale = args.max_side / max(w, h)
            im = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
        rel = f"images/{sku}.jpg"
        dest = img_dir / f"{sku}.jpg"
        im.save(dest, "JPEG", quality=args.quality, optimize=True)
        products.append(
            {
                "id": sku,
                "sku": sku,
                "name": name,
                "spec": clean_text(ws.cell(r, 4).value),
                "description": clean_text(ws.cell(r, 5).value),
                "category": clean_text(ws.cell(r, 6).value),
                "unit": clean_text(ws.cell(r, 7).value),
                "image": rel,
                "source": "seed",
            }
        )
        print(f"OK {sku} ({dest.stat().st_size} bytes)")

    catalog = {"version": 1, "products": products}
    (out / "catalog.json").write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"wrote {len(products)} products -> {out / 'catalog.json'}")


if __name__ == "__main__":
    main()
